# python/create_template.py
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

# Заголовок
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('ДОГОВОР ОКАЗАНИЯ УСЛУГ')
run.bold = True
run.font.size = Pt(16)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('№ {{contract_number}}')
run.font.size = Pt(14)

doc.add_paragraph('')

# Город и дата
p = doc.add_paragraph()
p.add_run('г. {{city}}').bold = False
p.add_run('\t\t\t\t\t')
p.add_run('«{{day}}» {{month}} {{year}} г.')

doc.add_paragraph('')

# Стороны
doc.add_paragraph(
    '{{company1_name}}, именуемое в дальнейшем «Заказчик», '
    'в лице {{company1_director}}, действующего на основании Устава, с одной стороны, и'
)
doc.add_paragraph(
    '{{company2_name}}, именуемое в дальнейшем «Исполнитель», '
    'в лице {{company2_director}}, действующего на основании Устава, с другой стороны,'
)
doc.add_paragraph('заключили настоящий Договор о нижеследующем:')

doc.add_paragraph('')

# Разделы
sections = [
    ('1. ПРЕДМЕТ ДОГОВОРА', [
        '1.1. Исполнитель обязуется оказать Заказчику услуги, указанные в п. 1.2 настоящего Договора, а Заказчик обязуется оплатить эти услуги.',
        '1.2. {{subject_description}}',
        '1.3. Услуги оказываются в соответствии с требованиями, согласованными Сторонами.'
    ]),
    ('2. СТОИМОСТЬ И ПОРЯДОК РАСЧЁТОВ', [
        '2.1. Стоимость услуг по настоящему Договору составляет {{amount}} ({{amount_words}}) рублей.',
        '2.2. Оплата производится Заказчиком в течение {{payment_days}} рабочих дней с момента подписания Акта оказанных услуг.',
        '2.3. Оплата производится путём перечисления денежных средств на расчётный счёт Исполнителя.'
    ]),
    ('3. СРОКИ ОКАЗАНИЯ УСЛУГ', [
        '3.1. Исполнитель обязуется оказать услуги в срок до {{end_date}}.',
        '3.2. В случае невозможности исполнения в указанный срок, Исполнитель обязан уведомить Заказчика не позднее чем за {{notice_days}} рабочих дней.'
    ]),
    ('4. ОТВЕТСТВЕННОСТЬ СТОРОН', [
        '4.1. За неисполнение или ненадлежащее исполнение обязательств по настоящему Договору Стороны несут ответственность в соответствии с действующим законодательством РФ.',
        '4.2. В случае просрочки оплаты Заказчик уплачивает Исполнителю пени в размере {{penalty_percent}}% от неоплаченной суммы за каждый день просрочки.',
        '4.3. {{additional_conditions}}'
    ]),
    ('5. ФОРС-МАЖОР', [
        '5.1. Стороны освобождаются от ответственности за частичное или полное неисполнение обязательств по настоящему Договору, если оно явилось следствием обстоятельств непреодолимой силы.'
    ]),
    ('6. РАЗРЕШЕНИЕ СПОРОВ', [
        '6.1. Все споры и разногласия решаются путём переговоров.',
        '6.2. В случае невозможности разрешения споров путём переговоров, они подлежат рассмотрению в суде по месту нахождения {{court_location}}.'
    ]),
    ('7. СРОК ДЕЙСТВИЯ ДОГОВОРА', [
        '7.1. Настоящий Договор вступает в силу с момента его подписания и действует до {{end_date}}.',
        '7.2. Договор может быть расторгнут досрочно по соглашению Сторон.'
    ]),
    ('8. РЕКВИЗИТЫ И ПОДПИСИ СТОРОН', [
        'ЗАКАЗЧИК:\n{{company1_name}}\nИНН: {{company1_inn}}\nАдрес: {{company1_address}}\n\nПодпись: _________________ / {{company1_director}} /',
        'ИСПОЛНИТЕЛЬ:\n{{company2_name}}\nИНН: {{company2_inn}}\nАдрес: {{company2_address}}\n\nПодпись: _________________ / {{company2_director}} /'
    ])
]

for title, items in sections:
    h = doc.add_paragraph()
    run = h.add_run(title)
    run.bold = True
    for item in items:
        doc.add_paragraph(item)
    doc.add_paragraph('')

doc.save('python/templates/dogovor_uslugi.docx')
print('✅ Шаблон создан: python/templates/dogovor_uslugi.docx')