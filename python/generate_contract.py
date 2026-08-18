import sys
import io
import json
import os
from datetime import datetime
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from rag_engine import (
    VERSION, init_vector_db, detect_best_document, get_document_text,
    clean_and_correct_input, ask_smart_questions, normalize_type,
    generate_document_structure, polish_document, enforce_structure,
    answer_question, extract_text, summarize_document,
)


def create_word_from_structure(structure, path):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)

    for line in structure.get('header', []):
        line = str(line).strip()
        if not line:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(line)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)

    if structure.get('header'):
        doc.add_paragraph()

    title = str(structure.get('title', 'ДОКУМЕНТ')).strip()
    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(16)
        run.font.name = 'Times New Roman'

    for line in structure.get('body', []):
        line = str(line).strip()
        if not line:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Pt(28)
        run = p.add_run(line)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)

    doc.add_paragraph()

    for line in structure.get('signature', []):
        line = str(line).strip()
        if not line:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(line)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)

    doc.save(path)


def structure_to_text(structure):
    parts = ['\n'.join(str(l) for l in structure.get('header', [])), '',
             str(structure.get('title', '')).upper(), '',
             '\n\n'.join(str(l) for l in structure.get('body', [])), '',
             '\n'.join(str(l) for l in structure.get('signature', []))]
    return '\n'.join(parts)


def create_pdf(docx_path, pdf_path):
    """Пытается создать PDF тремя способами: docx2pdf → soffice → None"""
    
    # Способ 1: docx2pdf (работает через установленный MS Word на Windows)
    try:
        from docx2pdf import convert
        convert(docx_path, pdf_path)
        if os.path.exists(pdf_path):
            print("PDF created via docx2pdf", file=sys.stderr)
            return True
    except ImportError:
        print("docx2pdf not installed, trying soffice...", file=sys.stderr)
    except Exception as e:
        print("docx2pdf error: " + str(e), file=sys.stderr)

    # Способ 2: LibreOffice (soffice)
    try:
        import subprocess
        out_dir = os.path.dirname(pdf_path)
        soffice = 'soffice'
        if sys.platform == 'win32':
            for prog in [os.environ.get('PROGRAMFILES', ''), os.environ.get('PROGRAMFILES(X86)', '')]:
                candidate = os.path.join(prog, 'LibreOffice', 'program', 'soffice.exe')
                if os.path.exists(candidate):
                    soffice = candidate
                    break
        
        result = subprocess.run(
            [soffice, '--headless', '--convert-to', 'pdf', '--outdir', out_dir, docx_path],
            capture_output=True, timeout=60
        )
        if result.returncode == 0:
            base = os.path.splitext(os.path.basename(docx_path))[0]
            generated = os.path.join(out_dir, base + '.pdf')
            if os.path.exists(generated) and generated != pdf_path:
                os.replace(generated, pdf_path)
            if os.path.exists(pdf_path):
                print("PDF created via soffice", file=sys.stderr)
                return True
    except Exception as e:
        print("soffice error: " + str(e), file=sys.stderr)

    # Способ 3: не получилось — вернём False, PHP создаст через dompdf
    print("PDF not created locally, PHP will handle it", file=sys.stderr)
    return False


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"success": False, "error": "No input"}, ensure_ascii=False))
        return

    try:
        with open(sys.argv[1], encoding='utf-8') as f:
            data = json.load(f)
    except Exception as e:
        print(json.dumps({"success": False, "error": "JSON: " + str(e)}, ensure_ascii=False))
        return

    try:
        mode = data.get('mode', 'generate')
        log_msg = "ENGINE v" + VERSION + " | mode=" + mode
        print(log_msg, file=sys.stderr)
        init_vector_db()

        if mode == 'chat':
            ans = answer_question(data.get('filename', ''), data.get('question', ''))
            print(json.dumps({"success": True, "answer": ans, "engine": VERSION}, ensure_ascii=False))
            return

        if mode == 'summarize':
            try:
                text = extract_text(data.get('filepath', ''))
            except Exception as e:
                print(json.dumps({"success": False, "error": str(e)}, ensure_ascii=False))
                return
            if not text:
                print(json.dumps({"success": False, "error": 'Не смог прочитать файл'}, ensure_ascii=False))
                return
            print(json.dumps({"success": True, "summary": summarize_document(text), "engine": VERSION}, ensure_ascii=False))
            return

        # ШАГ 1: чистим анкету
        print("STEP 1: cleaning...", file=sys.stderr)
        cleaned = clean_and_correct_input(data)
        fixed_title = normalize_type(data.get('doc_type'))
        cleaned['doc_type'] = fixed_title
        
        # 📎 Если пользователь прикрепил файл — добавляем в данные (для ИИ как контекст)
        if data.get('attachment_path'):
            cleaned['attachment_path'] = data['attachment_path']
            print("Attachment provided: " + data['attachment_path'], file=sys.stderr)
        
        print("Cleaned: " + json.dumps(cleaned, ensure_ascii=False), file=sys.stderr)

        # РЕЖИМ ВОПРОСОВ
        if mode == 'analyze':
            questions = ask_smart_questions(cleaned)
            best_file = detect_best_document(fixed_title + ' ' + cleaned.get('subject', ''))
            print(json.dumps({
                "success": True, "questions": questions,
                "sample": best_file or '', "engine": VERSION,
            }, ensure_ascii=False))
            return

        # РЕЖИМ ГЕНЕРАЦИИ
        best_file = detect_best_document(fixed_title + ' ' + cleaned.get('subject', ''))
        sample_text = get_document_text(best_file) if best_file else ''
        source_name = best_file or 'Составлено без образца'
        print("STEP 2: sample = " + source_name, file=sys.stderr)

        print("STEP 3: generating...", file=sys.stderr)
        structure = generate_document_structure(cleaned, sample_text, fixed_title)

        body_words = sum(len(str(b).split()) for b in structure.get('body', []))
        if body_words < 60:
            print("Too short (" + str(body_words) + " words), retrying...", file=sys.stderr)
            structure = generate_document_structure(
                cleaned, sample_text, fixed_title,
                retry_note="ПРОШЛАЯ ПОПЫТКА БЫЛА СЛИШКОМ КОРОТКОЙ. Пиши 120–250 слов!")

        print("STEP 4: polishing...", file=sys.stderr)
        structure = polish_document(structure)
        structure = enforce_structure(structure, cleaned, fixed_title)

        out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'output')
        os.makedirs(out_dir, exist_ok=True)
        base_name = 'new_' + datetime.now().strftime('%Y%m%d_%H%M%S')
        docx_path = os.path.join(out_dir, base_name + '.docx')
        pdf_path = os.path.join(out_dir, base_name + '.pdf')

        create_word_from_structure(structure, docx_path)
        
        # 📕 Генерируем PDF только если выбран PDF или если хотим всегда иметь PDF
        output_format = data.get('output_format', 'word')
        pdf_created = False
        
        if output_format == 'pdf' or True:  # Можно убрать `or True` если нужен PDF только когда выбран
            pdf_created = create_pdf(docx_path, pdf_path)
            if pdf_created:
                print("PDF successfully created", file=sys.stderr)
            else:
                print("PDF creation deferred to PHP (dompdf)", file=sys.stderr)

        print(json.dumps({
            "success": True,
            "filename": base_name + '.docx',
            "filepath": docx_path.replace('\\', '/'),
            "pdf_filepath": pdf_path.replace('\\', '/') if pdf_created else None,
            "detected_type": source_name,
            "detected_filename": source_name,
            "text": structure_to_text(structure),
            "engine": VERSION,
        }, ensure_ascii=False))

    except Exception as e:
        import traceback
        traceback.print_exc(file=sys.stderr)
        print(json.dumps({"success": False, "error": str(e)}, ensure_ascii=False))


if __name__ == '__main__':
    main()